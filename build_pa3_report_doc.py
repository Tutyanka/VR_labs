from pathlib import Path
import textwrap
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "PA3_Zvit_Ukr.docx"
ASSETS = ROOT / "report_assets"
ASSETS.mkdir(exist_ok=True)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_border(cell, color="DADCE0", size="8"):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = f"w:{edge}"
        element = tc_borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tc_borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def code_border(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.first_child_found_in("w:pBdr")
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "8")
        el.set(qn("w:space"), "6")
        el.set(qn("w:color"), "DADCE0")
        p_bdr.append(el)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F6F8FA")
    p_pr.append(shd)


def apply_run_font(run, name="Arial", size=11, bold=False, color="000000", italic=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def style_paragraph(paragraph, *, after=8, before=0, line=1.15, align=WD_ALIGN_PARAGRAPH.LEFT):
    pf = paragraph.paragraph_format
    pf.space_after = Pt(after)
    pf.space_before = Pt(before)
    pf.line_spacing = line
    paragraph.alignment = align


def add_body(doc, text):
    p = doc.add_paragraph()
    style_paragraph(p, after=8, line=1.15)
    run = p.add_run(text)
    apply_run_font(run, size=11)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    sizes = {1: 20, 2: 16, 3: 14}
    colors = {1: "000000", 2: "000000", 3: "434343"}
    before = {1: 18, 2: 14, 3: 10}
    after = {1: 6, 2: 4, 3: 3}
    style_paragraph(p, before=before[level], after=after[level], line=1.0)
    run = p.add_run(text)
    apply_run_font(run, size=sizes[level], color=colors[level], bold=False)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    style_paragraph(p, after=4, line=1.15)
    run = p.add_run(text)
    apply_run_font(run, size=11)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    style_paragraph(p, after=4, line=1.15)
    run = p.add_run(text)
    apply_run_font(run, size=11)
    return p


def add_code_block(doc, code):
    p = doc.add_paragraph()
    style_paragraph(p, after=10, line=1.1)
    code_border(p)
    lines = code.splitlines()
    for i, line in enumerate(lines):
        run = p.add_run(line)
        apply_run_font(run, name="Consolas", size=9)
        if i != len(lines) - 1:
            run.add_break()
    return p


def add_caption(doc, text):
    p = doc.add_paragraph()
    style_paragraph(p, after=10, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
    run = p.add_run(text)
    apply_run_font(run, size=10, italic=True, color="555555")
    return p


def extract_function(name):
    src = (ROOT / "index.html").read_text(encoding="utf-8")
    idx = src.find(name)
    if idx == -1:
        return f"// Функцію не знайдено: {name}"
    brace = src.find("{", idx)
    depth = 0
    end = brace
    while end < len(src):
        ch = src[end]
        if ch == "{":
            depth += 1
        elif ch == "}":
            depth -= 1
            if depth == 0:
                end += 1
                break
        end += 1
    return src[idx:end].replace("      ", "")


def make_placeholder(path, title, subtitle):
    width, height = 1400, 840
    image = Image.new("RGB", (width, height), "#EEF4FF")
    draw = ImageDraw.Draw(image)
    try:
        title_font = ImageFont.truetype("arial.ttf", 54)
        text_font = ImageFont.truetype("arial.ttf", 28)
        small_font = ImageFont.truetype("arial.ttf", 22)
    except OSError:
        title_font = ImageFont.load_default()
        text_font = ImageFont.load_default()
        small_font = ImageFont.load_default()

    draw.rounded_rectangle((70, 60, width - 70, height - 60), radius=28, fill="#FFFFFF", outline="#B9C8E5", width=4)
    draw.rounded_rectangle((110, 110, width - 110, 250), radius=18, fill="#F5F9FF", outline="#D4E0F7", width=2)
    draw.text((150, 145), "Ding-dong ПР3 Просторовий звук", fill="#16355F", font=title_font)
    draw.text((150, 225), title, fill="#2A5CAA", font=text_font)
    draw.rounded_rectangle((140, 310, width - 140, height - 180), radius=22, fill="#F8FBFF", outline="#C8D7F0", width=3)
    draw.ellipse((250, 390, 370, 510), fill="#17C7E7", outline="#0B8AA1", width=4)
    draw.ellipse((570, 360, 940, 660), outline="#79A8E8", width=5)
    draw.line((780, 510, 310, 450), fill="#7B8BA8", width=5)
    draw.text((170, 690), subtitle, fill="#4A5A75", font=text_font)
    note = "Тимчасовий скріншот для звіту. За потреби заміни його реальним знімком із http://localhost:8090 перед остаточною здачею."
    wrapped = textwrap.fill(note, width=82)
    draw.multiline_text((170, 730), wrapped, fill="#6B7280", font=small_font, spacing=6)
    image.save(path)


def add_screenshot_grid(doc, paths, captions):
    table = doc.add_table(rows=2, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [Inches(3.15), Inches(3.15)]
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = widths[idx]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cell)
            set_cell_shading(cell, "FFFFFF")
    items = list(zip(paths, captions))
    for idx, (path, caption) in enumerate(items):
        cell = table.rows[idx // 2].cells[idx % 2]
        cell_p = cell.paragraphs[0]
        cell_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = cell_p.add_run()
        run.add_picture(str(path), width=Inches(2.75))
        cap = cell.add_paragraph()
        style_paragraph(cap, after=4, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
        run = cap.add_run(caption)
        apply_run_font(run, size=9, italic=True, color="555555")


def set_page(section):
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)


def add_footer(section, text):
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(footer, after=0, line=1.0)
    run = footer.add_run(text)
    apply_run_font(run, size=9, color="666666")


def build_doc():
    for idx, (title, subtitle) in enumerate([
        ("Скріншот 1. Головне вікно з маркером джерела звуку", "На першому скріншоті має бути видно поверхню Ding-dong і блакитну сферу джерела звуку."),
        ("Скріншот 2. Керування spatial audio та фільтром", "Другий скріншот має показувати HUD із завантаженим треком, кнопкою відтворення і параметрами фільтра."),
        ("Скріншот 3. Підключення та калібрування телефона", "Третій скріншот має демонструвати активне сенсорне підключення і стан після калібрування."),
        ("Скріншот 4. Інше положення джерела після обертання", "Четвертий скріншот має показувати, як джерело звуку змінило положення після повороту телефона."),
    ], start=1):
        make_placeholder(ASSETS / f"screenshot_{idx}.png", title, subtitle)

    doc = Document()
    set_page(doc.sections[0])
    add_footer(doc.sections[0], "Звіт до ПР №3: Spatial Audio")

    for _ in range(6):
        add_body(doc, "")
    p = doc.add_paragraph()
    style_paragraph(p, after=6, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run("ПРАКТИЧНА РОБОТА №3")
    apply_run_font(r, size=20, bold=True)
    p = doc.add_paragraph()
    style_paragraph(p, after=10, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run("Просторовий звук для поверхні Ding-dong")
    apply_run_font(r, size=26)
    p = doc.add_paragraph()
    style_paragraph(p, after=20, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run("Комп'ютерна графіка для веб")
    apply_run_font(r, size=14, color="555555")
    for label, value in [
        ("Студент", "__________________"),
        ("Група", "__________________"),
        ("Варіант", "Фільтр високих частот"),
        ("Гілка репозиторію", "CGW"),
        ("Рік", "2026"),
    ]:
        p = doc.add_paragraph()
        style_paragraph(p, after=6, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
        r1 = p.add_run(f"{label}: ")
        apply_run_font(r1, size=12, bold=True)
        r2 = p.add_run(value)
        apply_run_font(r2, size=12)
    doc.add_page_break()

    add_heading(doc, "1. Опис завдання", 1)
    add_body(doc, "Метою цієї практичної роботи є розширення попереднього вебзастосунку з візуалізацією поверхні Ding-dong шляхом додавання підсистеми просторового звуку на основі HTML5 WebAudio API. Нова версія застосунку повинна використовувати напрацювання з практичної роботи №2, зокрема вже реалізовану WebGL-сцену, стереовідображення, текстурування, а також матеріальний інтерфейс керування на основі даних зі смартфона.")
    add_body(doc, "Головна функціональна вимога полягає в тому, що сама поверхня залишається нерухомою, а навколо її геометричного центру переміщується джерело звуку. Положення цього джерела повинно керуватися користувачем за допомогою матеріального інтерфейсу, тобто фізичного обертання телефона, сенсорні дані з якого передаються до браузера через локальний проміжний сервер і WebSocket-з'єднання.")
    add_body(doc, "Щоб користувач міг наочно спостерігати за роботою системи, просторове джерело звуку потрібно візуалізувати окремою сферою. У межах виконаної роботи ця сфера має яскравий блакитний колір і рухається в тих самих координатах, у яких у WebAudio розташовано джерело відтворення пісні.")
    add_body(doc, "Крім просторового розміщення звуку, завдання вимагає використання фільтра на основі `BiquadFilterNode`. Для цього варіанта реалізовується фільтр високих частот. Інтерфейс має містити чекбокс для ввімкнення або вимкнення фільтра, а також параметри для налаштування частоти зрізу і добротності `Q`.")
    add_body(doc, "Отже, практична робота поєднує повторне використання коду з ПР №2, тривимірну візуалізацію, сенсорне керування, просторовий звук і цифрову аудіообробку в одному інтерактивному вебпроєкті.")
    doc.add_page_break()

    add_heading(doc, "2. Теоретичні відомості", 1)
    add_heading(doc, "2.1 Просторовий звук у WebAudio", 2)
    add_body(doc, "Просторовий звук — це спосіб відтворення аудіосигналу, за якого слухач сприймає його положення в просторі: ліворуч чи праворуч, ближче чи далі, попереду чи позаду. У WebAudio API за таку поведінку відповідає `PannerNode`, який обробляє звуковий сигнал відповідно до координат джерела та координат слухача.")
    add_body(doc, "У реалізованому застосунку використовується модель панорамування `HRTF`. Абревіатура HRTF означає функцію, пов'язану з впливом голови та вух на сприйняття звуку. Вона наближено моделює те, як змінюється сигнал, перш ніж потрапити в лівий і правий канали. Завдяки цьому звук сприймається об'ємніше й переконливіше, ніж при звичайному стереобалансі.")
    add_body(doc, "Крім моделі панорамування, велике значення має модель затухання за відстанню. У цьому проєкті використовується обернена модель відстані, тобто гучність зменшується зі збільшенням відстані до джерела. Такий підхід краще узгоджується з реальним слуховим досвідом користувача.")
    add_heading(doc, "2.2 Положення слухача і джерела", 2)
    add_body(doc, "Для коректної роботи просторового звуку потрібно постійно оновлювати дві сутності: положення слухача і положення джерела звуку. У цій роботі слухач пов'язаний із віртуальною камерою сцени, тому при зміні кута огляду або положення камери змінюється і слухова перспектива.")
    add_body(doc, "Положення джерела обчислюється або зі смартфона, або з ручних слайдерів резервного керування. Коли ввімкнено телефонне керування, базовий вектор положення обертається каліброваною матрицею орієнтації, отриманою із сенсорних даних. Завдяки цьому рухається саме джерело звуку, а не сама поверхня.")
    add_heading(doc, "2.3 Tangible user interface", 2)
    add_body(doc, "Матеріальний інтерфейс керування — це спосіб взаємодії, коли фізичний об'єкт або реальний рух прямо впливає на цифрову систему. У цій роботі смартфон виконує роль такого фізичного контролера. Його гіроскопічні або кватерніонні дані надходять на локальний проміжний сервер, а звідти передаються на вебсторінку через WebSocket-з'єднання.")
    doc.add_page_break()

    add_heading(doc, "2.4 Кватерніони та калібрування", 2)
    add_body(doc, "Для опису орієнтації в просторі доцільно використовувати кватерніони, оскільки вони стійкіші до накопичення похибок і не мають таких проблем, як кути Ейлера. Отриманий кватерніон перетворюється на матрицю обертання, яка далі використовується для обертання вектора положення звукового джерела.")
    add_body(doc, "Оскільки телефон може починати роботу в довільному положенні, у застосунку реалізовано калібрування. Воно зберігає обернену матрицю поточного стану й надалі вважає це положення нейтральним. Завдяки цьому подальші повороти телефона інтерпретуються відносно зручної для користувача стартової пози.")
    add_heading(doc, "2.5 Фільтр високих частот", 2)
    add_body(doc, "Фільтр високих частот пропускає верхні частоти і послаблює нижні. У WebAudio він створюється за допомогою `BiquadFilterNode`, якщо для параметра `type` встановити значення `highpass`. Такий фільтр дає змогу змінити тембр композиції, зробивши її звучання світлішим, тоншим або менш насиченим басами.")
    add_body(doc, "Найважливішими параметрами є частота зрізу та добротність `Q`. Частота зрізу визначає межу, нижче якої спектральні компоненти пригнічуються, а `Q` впливає на різкість переходу в області зрізу. У застосунку ці параметри винесені в окремі елементи керування, що дає змогу слухачеві відразу чути результат зміни налаштувань.")
    add_body(doc, "Поєднання просторового переміщення звуку і спектральної обробки показує, що сучасний браузер може одночасно працювати і з геометрією звучання, і з цифровою аудіофільтрацією в реальному часі.")
    doc.add_page_break()

    add_heading(doc, "3. Деталі реалізації", 1)
    add_heading(doc, "3.1 Повторне використання ПР №2", 2)
    add_body(doc, "Під час виконання цієї роботи було збережено основну архітектуру попереднього застосунку: генерацію поверхні Ding-dong, WebGL-відображення, стереокамеру, систему матеріалів і текстур, вебкамерний фон, а також Node.js-сервер-посередник для приймання сенсорних даних зі смартфона.")
    add_body(doc, "Найважливіша зміна стосується логіки застосування сенсорної орієнтації. Якщо в ПР №2 обертання телефона змінювало орієнтацію моделі, то в ПР №3 ці самі дані використовуються для обчислення орбіти джерела звуку навколо центру поверхні.")
    add_heading(doc, "3.2 Побудова аудіографа", 2)
    add_body(doc, "Аудіопідсистема створюється не одразу при завантаженні сторінки, а під час першого запуску відтворення. Це дозволяє уникнути проблем з автоматичним блокуванням аудіо в браузері. Після натискання кнопки `play` формується граф такого вигляду: `HTMLAudioElement -> MediaElementSourceNode -> BiquadFilterNode (за потреби) -> PannerNode -> destination`.")
    add_bullet(doc, "Для панорамування використовується `panningModel = \"HRTF\"`.")
    add_bullet(doc, "Для затухання за відстанню використовується `distanceModel = \"inverse\"`.")
    add_bullet(doc, "Параметри `refDistance`, `maxDistance` і `rolloffFactor` підібрано для зручного демо-ефекту.")
    add_bullet(doc, "Фільтр реалізовано через `BiquadFilterNode` типу `highpass`.")
    add_heading(doc, "3.3 Обчислення положення джерела", 2)
    add_body(doc, "У застосунку передбачено два режими роботи джерела. Перший — ручний режим, у якому користувач змінює азимут та кут підйому через повзунки. Другий — сенсорний режим, у якому положення джерела формується обертанням вектора `[radius, 0, height]` за допомогою матриці, обчисленої з орієнтації смартфона.")
    doc.add_page_break()

    add_heading(doc, "3.4 Синхронізація графіки та звуку", 2)
    add_body(doc, "Вектор `soundSourcePos` використовується відразу у двох місцях: як координати для `PannerNode` і як координати для відображення блакитної сфери у WebGL-сцені. Завдяки цьому візуальна і слухова частини застосунку повністю узгоджені.")
    add_body(doc, "Також перед кожним кадром оновлюється положення і напрямок слухача на основі поточної камери сцени. Це необхідно для того, щоб звук змінювався разом із віртуальною точкою огляду, а не залишався статичним відносно екрана.")
    add_heading(doc, "3.5 Нові елементи інтерфейсу", 2)
    add_body(doc, "У HUD були додані нові елементи: кнопка запуску й паузи відтворення, вибір локального аудіофайлу у форматі `.mp3` або `.ogg`, повзунки радіуса і висоти орбіти, повзунки ручного азимута та кута підйому, чекбокс увімкнення фільтра, слайдер частоти зрізу, слайдер `Q`, а також текстова індикація поточних координат джерела.")
    add_body(doc, "Ці елементи дають можливість не тільки продемонструвати виконання вимог завдання, а й зручно пояснити реалізацію під час захисту. Користувач може підключити телефон, відкалібрувати його, завантажити улюблену композицію і в реальному часі спостерігати за рухом джерела та зміною тембру.")
    add_heading(doc, "3.6 Практичні особливості", 2)
    add_body(doc, "Для повноцінного сприйняття просторового ефекту бажано використовувати навушники. Локальний сервер запускається командою `node bridge-server.js 8090`, після чого сторінка відкривається за адресою `http://localhost:8090`, а телефон надсилає сенсорні дані на маршрут `/sensor`.")
    add_body(doc, "У підсумку було реалізовано цілісний браузерний застосунок, який поєднує WebGL-графіку, tangible control, просторовий звук і цифрову аудіообробку без використання зовнішніх фронтенд-фреймворків.")
    doc.add_page_break()

    add_heading(doc, "4. Інструкція користувача зі скріншотами", 1)
    add_heading(doc, "4.1 Запуск застосунку", 2)
    for item in [
        "Відкрити папку проєкту у локальному середовищі.",
        "У терміналі виконати команду `node bridge-server.js 8090`.",
        "Відкрити у браузері адресу `http://localhost:8090`.",
        "У блоці керування просторовим звуком вибрати локальний аудіофайл `.mp3` або `.ogg`.",
        "Натиснути кнопку `play`, щоб запустити відтворення композиції.",
    ]:
        add_number(doc, item)
    add_heading(doc, "4.2 Підключення телефона", 2)
    for item in [
        "Підключити комп'ютер і телефон до однієї мережі Wi-Fi.",
        "Налаштувати сенсорний застосунок на телефоні на надсилання даних до `http://<ip-комп'ютера>:8090/sensor`.",
        "У вебсторінці вказати WebSocket-адресу `ws://<ip-комп'ютера>:8090`, якщо потрібно працювати не через localhost.",
        "Натиснути `connect`, а потім, тримаючи телефон у зручному нейтральному положенні, натиснути `calibrate`.",
        "Переконатися, що прапорець `Use phone for sound orbit` увімкнено.",
    ]:
        add_number(doc, item)
    add_screenshot_grid(
        doc,
        [ASSETS / "screenshot_1.png", ASSETS / "screenshot_2.png", ASSETS / "screenshot_3.png", ASSETS / "screenshot_4.png"],
        [
            "Рисунок 1. Головне вікно застосунку з маркером джерела звуку.",
            "Рисунок 2. Панель керування просторовим звуком і фільтром.",
            "Рисунок 3. Стан підключення телефона та калібрування.",
            "Рисунок 4. Інше положення джерела після обертання телефона."
        ],
    )
    add_caption(doc, "Рисунки 1–4. Тимчасові зображення для звіту, які можна замінити реальними скріншотами застосунку.")
    doc.add_page_break()

    add_heading(doc, "4.3 Використання фільтра та демонстрація", 2)
    add_body(doc, "Після запуску відтворення можна ввімкнути прапорець фільтра високих частот. Далі варто змінювати частоту зрізу, щоб послаблювати або сильніше вирізати низькі частоти, а також регулювати `Q`, щоб змінювати різкість переходу в області зрізу. Найзручніше демонструвати цю можливість, порівнюючи звучання пісні до ввімкнення фільтра і після нього.")
    add_body(doc, "Під час демонстрації важливо звернути увагу на блакитну сферу, оскільки саме вона показує, де в просторі зараз розташоване джерело. Якщо сенсорне підключення тимчасово недоступне, просторову поведінку можна все одно показати через ручні слайдери азимута та кута підйому.")
    add_heading(doc, "4.4 Рекомендована послідовність показу", 2)
    for item in [
        "Показати поверхню Ding-dong і блакитну сферу джерела звуку.",
        "Завантажити улюблену композицію і ввімкнути її відтворення.",
        "Повернути телефон і пояснити, що рухається саме джерело звуку, а не поверхня.",
        "Увімкнути фільтр високих частот і змінити параметри зрізу та `Q`.",
        "Трохи змінити камеру, щоб показати узгодження між зображенням і просторовим звучанням.",
    ]:
        add_number(doc, item)
    add_body(doc, "Попередня сторінка містить місця для чотирьох скріншотів. Перед остаточною здачею їх варто замінити реальними знімками з локального застосунку, щоб звіт повністю відповідав формальним вимогам завдання.")
    add_body(doc, "Отже, цей розділ виконує подвійну роль: короткої інструкції для користувача і сценарію для захисту та демонстрації роботи.")
    doc.add_page_break()

    add_heading(doc, "5. Зразок вихідного коду", 1)
    add_heading(doc, "5.1 Обчислення положення джерела та оновлення слухача", 2)
    add_body(doc, "Наведений нижче фрагмент коду показує, як обчислюється положення джерела звуку на основі або каліброваної орієнтації телефона, або ручних слайдерів резервного керування. Ця ж частина коду передає координати в `PannerNode` і синхронізує слухача з поточною камерою сцени.")
    add_code_block(doc, extract_function("function computeSoundSourcePosition()"))
    add_code_block(doc, extract_function("function updateAudioScene(cameraBasis)"))
    add_body(doc, "Саме цей фрагмент реалізує головну вимогу завдання: джерело обертається навколо центру поверхні, а сама поверхня залишається нерухомою. Тут також видно зв'язок між візуальною камерою і слуховою перспективою.")
    doc.add_page_break()

    add_heading(doc, "5.2 Аудіограф і фільтр високих частот", 2)
    add_body(doc, "Наступний фрагмент демонструє створення аудіографа. Спочатку формується `MediaElementSourceNode`, потім за потреби застосовується `BiquadFilterNode`, а результат передається до `PannerNode`. Остаточна схема з'єднань залежить від того, чи ввімкнено чекбокс фільтра.")
    add_code_block(doc, extract_function("function ensureAudioGraph()"))
    add_code_block(doc, extract_function("function updateFilterParams()"))
    add_code_block(doc, extract_function("function rebuildAudioRouting()"))
    add_body(doc, "Разом ці функції показують ключову реалізацію варіантного завдання: побудову просторового аудіо і керування фільтром високих частот у браузері в реальному часі.")

    doc.save(OUTPUT)


if __name__ == "__main__":
    build_doc()
    print(OUTPUT)
